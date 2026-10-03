"""Laporan HR/manajer: tren partisipasi membaca per divisi (fungsi) per minggu.

Hanya agregat. Fungsi dengan anggota lebih sedikit dari `reports.min_group_size` digabung
menjadi satu baris agar angka tidak bisa ditelusuri ke individu.
"""

import csv
import io
from collections import defaultdict
from datetime import date, timedelta

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from pymongo.asynchronous.database import AsyncDatabase

from app.core import clock
from app.repositories import catalog

SMALL_GROUP = "Fungsi kecil (digabung)"
DEFAULT_MIN_GROUP = 3


def _week_label(start: date) -> str:
    iso = start.isocalendar()
    return f"{iso.year}-W{iso.week:02d}"


async def participation(db: AsyncDatabase, weeks: int) -> dict:
    tz = str(await catalog.get_setting(db, "team.timezone", "Asia/Jakarta"))
    min_group = int(await catalog.get_setting(db, "reports.min_group_size", DEFAULT_MIN_GROUP))
    today = date.fromisoformat(clock.local_date(clock.now(), tz))
    this_monday = today - timedelta(days=today.weekday())
    starts = [this_monday - timedelta(weeks=weeks - 1 - i) for i in range(weeks)]
    first, last = starts[0], starts[-1] + timedelta(days=6)

    functions = {f["_id"]: f["name"] async for f in db["functions"].find({}, {"name": 1})}
    users = (
        await db["users"]
        .find(
            {"status": "active", "function_id": {"$ne": None}}, {"function_id": 1, "created_at": 1}
        )
        .to_list()
    )
    size = defaultdict(int)
    for u in users:
        size[u["function_id"]] += 1
    group_of = {
        u["_id"]: functions.get(u["function_id"], SMALL_GROUP)
        if size[u["function_id"]] >= min_group
        else SMALL_GROUP
        for u in users
    }

    sessions = (
        await db["reading_sessions"]
        .find(
            {
                "status": "completed",
                "local_date": {"$gte": first.isoformat(), "$lte": last.isoformat()},
                "user_id": {"$in": list(group_of)},
            },
            {"user_id": 1, "local_date": 1, "active_seconds": 1},
        )
        .to_list()
    )

    # (grup, minggu) → pembaca unik & detik
    readers: dict = defaultdict(set)
    seconds: dict = defaultdict(int)
    for s in sessions:
        day = date.fromisoformat(s["local_date"])
        week = (day - first).days // 7
        key = (group_of[s["user_id"]], week)
        readers[key].add(s["user_id"])
        seconds[key] += s["active_seconds"]

    def members_at(group: str, week: int) -> int:
        end = starts[week] + timedelta(days=7)
        return sum(
            1
            for u in users
            if group_of[u["_id"]] == group
            and (not u.get("created_at") or u["created_at"].date() < end)
        )

    groups = sorted({g for g in group_of.values()}, key=lambda g: (g == SMALL_GROUP, g))
    rows = []
    for group in groups:
        series = []
        for w in range(weeks):
            members = members_at(group, w)
            count = len(readers[(group, w)])
            minutes = round(seconds[(group, w)] / 60)
            series.append(
                {
                    "week": _week_label(starts[w]),
                    "members": members,
                    "readers": count,
                    "rate": round(count / members, 3) if members else 0,
                    "minutes": minutes,
                    "avg_minutes_per_member": round(minutes / members, 1) if members else 0,
                }
            )
        rows.append(
            {
                "function": group,
                "members": sum(1 for g in group_of.values() if g == group),
                "series": series,
            }
        )

    overall = []
    for w in range(weeks):
        members = sum(members_at(g, w) for g in groups)
        count = sum(len(readers[(g, w)]) for g in groups)
        minutes = round(sum(seconds[(g, w)] for g in groups) / 60)
        overall.append(
            {
                "week": _week_label(starts[w]),
                "members": members,
                "readers": count,
                "rate": round(count / members, 3) if members else 0,
                "minutes": minutes,
                "avg_minutes_per_member": round(minutes / members, 1) if members else 0,
            }
        )
    return {
        "weeks": [_week_label(s) for s in starts],
        "start": first.isoformat(),
        "end": last.isoformat(),
        "min_group_size": min_group,
        "functions": rows,
        "overall": overall,
    }


HEADERS = ["Fungsi", "Minggu", "Anggota", "Pembaca", "Partisipasi (%)", "Menit", "Menit/anggota"]


def _flat(report: dict) -> list[list]:
    out = []
    for row in [*report["functions"], {"function": "Semua fungsi", "series": report["overall"]}]:
        for p in row["series"]:
            out.append(
                [
                    row["function"],
                    p["week"],
                    p["members"],
                    p["readers"],
                    round(p["rate"] * 100, 1),
                    p["minutes"],
                    p["avg_minutes_per_member"],
                ]
            )
    return out


def to_csv(report: dict) -> bytes:
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow(HEADERS)
    writer.writerows(_flat(report))
    return ("﻿" + buf.getvalue()).encode("utf-8")  # BOM agar Excel membaca UTF-8


def to_xlsx(report: dict) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Partisipasi per divisi"
    ws.append(HEADERS)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="2563EB")
    for line in _flat(report):
        ws.append(line)
    for col, width in zip("ABCDEFG", (28, 12, 10, 10, 16, 10, 14), strict=True):
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "A2"
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
