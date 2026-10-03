"""Export laporan Admin ke Excel (openpyxl) dan PDF (fpdf2)."""

import io
from pathlib import Path

from fpdf import FPDF
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

FONT_DIR = Path(__file__).resolve().parent.parent / "assets" / "fonts"
HEADER_FILL = PatternFill("solid", fgColor="6C4DF6")
HEADER_FONT = Font(bold=True, color="FFFFFF")


def _sheet(wb: Workbook, title: str, headers: list[str], rows: list[list], first: bool = False):
    ws = wb.active if first else wb.create_sheet()
    ws.title = title
    ws.append(headers)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical="center")
    for row in rows:
        ws.append(row)
    for i, header in enumerate(headers, start=1):
        width = max([len(str(header))] + [len(str(r[i - 1])) for r in rows] + [8])
        ws.column_dimensions[get_column_letter(i)].width = min(width + 2, 48)
    ws.freeze_panes = "A2"
    return ws


def summary_pairs(d: dict) -> list[list]:
    return [
        ["Periode", f"{d['start']} s.d. {d['end']} ({d['days']} hari)"],
        ["Anggota aktif", d["members"]],
        ["Pembaca aktif", d["active_readers"]],
        ["Tingkat partisipasi", f"{round(d['participation_rate'] * 100)}%"],
        ["Total menit baca", d["total_minutes"]],
        ["Rata-rata menit per pembaca", d["avg_minutes_per_reader"]],
        ["Rata-rata menit/hari per pembaca", d["avg_daily_minutes_per_reader"]],
        ["Sesi baca", d["sessions"]],
        ["Catatan", d["notes"]],
        ["Buku selesai", d["books_finished"]],
        *[
            [f"Status {k.replace('_', ' ').title()}", v]
            for k, v in d["authenticity_counts"].items()
        ],
    ]


def to_xlsx(dashboard: dict, users: list[dict]) -> bytes:
    wb = Workbook()
    _sheet(wb, "Ringkasan", ["Metrik", "Nilai"], summary_pairs(dashboard), first=True)
    _sheet(
        wb,
        "Harian",
        ["Tanggal", "Menit", "Pembaca"],
        [[d["date"], d["minutes"], d["readers"]] for d in dashboard["daily"]],
    )
    _sheet(
        wb,
        "Heatmap Fungsi",
        ["Fungsi", "Anggota", *dashboard["weekdays"]],
        [[h["function"], h["members"], *h["values"]] for h in dashboard["heatmap"]],
    )
    _sheet(
        wb,
        "Pengguna",
        ["Nama", "No. HP", "Fungsi", "Sesi", "Menit", "Catatan", "Poin", "Status Authenticity"],
        [
            [
                u["name"],
                u["phone"],
                u["function"],
                u["sessions"],
                u["minutes"],
                u["notes"],
                u["points"],
                u["status"],
            ]
            for u in users
        ],
    )
    _sheet(
        wb,
        "Observer",
        ["Nama", "Fungsi", "Catatan", "Komentar", "Reaksi", "Rasio"],
        [
            [
                o["name"],
                o["function"] or "-",
                o["own_notes"],
                o["comments_given"],
                o["likes_given"],
                o["contribution_ratio"],
            ]
            for o in dashboard["observers"]
        ],
    )
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


class _Report(FPDF):
    app_name = "ReadQuest"

    def header(self) -> None:
        self.set_font("Nunito", "B", 9)
        self.set_text_color(108, 77, 246)
        self.cell(0, 6, f"{self.app_name} · Laporan Admin", align="R")
        self.ln(8)

    def footer(self) -> None:
        self.set_y(-12)
        self.set_font("Nunito", "", 8)
        self.set_text_color(120, 120, 120)
        self.cell(0, 6, f"Halaman {self.page_no()}", align="C")


def _table(pdf: FPDF, headers: list[str], rows: list[list], widths: list[float]) -> None:
    pdf.set_font("Nunito", "B", 9)
    pdf.set_fill_color(108, 77, 246)
    pdf.set_text_color(255, 255, 255)
    for h, w in zip(headers, widths, strict=True):
        pdf.cell(w, 7, str(h), border=0, fill=True)
    pdf.ln()
    pdf.set_font("Nunito", "", 9)
    pdf.set_text_color(31, 27, 46)
    for i, row in enumerate(rows):
        fill = i % 2 == 1
        pdf.set_fill_color(246, 239, 230)
        for value, w in zip(row, widths, strict=True):
            text = str(value)
            while pdf.get_string_width(text) > w - 2 and len(text) > 1:
                text = text[:-2] + "…"
            pdf.cell(w, 6.5, text, fill=fill)
        pdf.ln()
    pdf.ln(4)


def to_pdf(dashboard: dict, users: list[dict], app_name: str = "ReadQuest") -> bytes:
    pdf = _Report(orientation="P", unit="mm", format="A4")
    pdf.app_name = app_name
    pdf.add_font("Nunito", "", str(FONT_DIR / "Nunito-Regular.ttf"))
    pdf.add_font("Nunito", "B", str(FONT_DIR / "Nunito-Bold.ttf"))
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()
    pdf.set_font("Nunito", "B", 18)
    pdf.set_text_color(31, 27, 46)
    pdf.cell(0, 10, "Laporan Membaca Tim")
    pdf.ln(12)

    pdf.set_font("Nunito", "B", 12)
    pdf.cell(0, 8, "Ringkasan")
    pdf.ln(9)
    _table(pdf, ["Metrik", "Nilai"], summary_pairs(dashboard), [90, 90])

    pdf.set_font("Nunito", "B", 12)
    pdf.cell(0, 8, "Rata-rata menit per anggota per hari (per fungsi)")
    pdf.ln(9)
    _table(
        pdf,
        ["Fungsi", "Agt", *dashboard["weekdays"]],
        [[h["function"], h["members"], *h["values"]] for h in dashboard["heatmap"]],
        [50, 12] + [17] * 7,
    )

    pdf.set_font("Nunito", "B", 12)
    pdf.cell(0, 8, "Rekap pengguna")
    pdf.ln(9)
    _table(
        pdf,
        ["Nama", "Fungsi", "Sesi", "Menit", "Catatan", "Poin", "Status"],
        [
            [
                u["name"],
                u["function"],
                u["sessions"],
                u["minutes"],
                u["notes"],
                u["points"],
                u["status"],
            ]
            for u in users
        ],
        [44, 38, 14, 16, 16, 16, 36],
    )
    return bytes(pdf.output())
