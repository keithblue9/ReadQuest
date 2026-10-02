import io
import uuid

from bson import ObjectId
from openpyxl import load_workbook

from tests.conftest import (
    USER_PIN,
    admin_headers,
    auth_header,
    create_book,
    me,
    onboarded_user,
    quick_post,
)


def note() -> str:
    return (
        f"Catatan {uuid.uuid4().hex}: penulis menjelaskan bagaimana kebiasaan kecil yang "
        "konsisten membentuk identitas baru. Saya mencatat tiga ide yang ingin dicoba bersama tim "
        "minggu ini, mulai dari menaruh buku di meja kerja sampai membaca sebelum rapat pagi."
    )


def code() -> str:
    return f"uji-{uuid.uuid4().hex[:8]}"


async def test_member_cannot_access_admin(client):
    member = await onboarded_user(client)
    for path in [
        "/api/v1/admin/dashboard",
        "/api/v1/admin/resources/badges",
        "/api/v1/admin/settings",
        "/api/v1/admin/users",
        "/api/v1/admin/moderation",
        "/api/v1/admin/audit",
        "/api/v1/admin/export.xlsx",
        "/api/v1/admin/lookups",
    ]:
        assert (await client.get(path, headers=member)).status_code == 403, path


async def test_dashboard_and_exports(client, fake_clock):
    admin = await admin_headers(client)
    reader, lurker = await onboarded_user(client), await onboarded_user(client)
    book = await create_book(client, reader)
    post = (await quick_post(client, reader, book["id"], note()))["post"]
    await client.put(f"/api/v1/posts/{post['id']}/reaction", json={"type": "like"}, headers=lurker)

    data = (await client.get("/api/v1/admin/dashboard", params={"days": 30}, headers=admin)).json()
    assert data["active_readers"] >= 1 and data["members"] >= 2
    assert len(data["daily"]) == 30 and data["weekdays"][0] == "Sen"
    assert {"function", "members", "values"} <= set(data["heatmap"][0])
    lurker_id = (await me(client, lurker))["id"]
    assert any(o["user_id"] == lurker_id for o in data["observers"])
    assert data["top_books"]

    xlsx = await client.get("/api/v1/admin/export.xlsx", headers=admin)
    assert xlsx.status_code == 200
    assert "attachment" in xlsx.headers["content-disposition"]
    wb = load_workbook(io.BytesIO(xlsx.content))
    assert wb.sheetnames == ["Ringkasan", "Harian", "Heatmap Fungsi", "Pengguna", "Observer"]
    pdf = await client.get("/api/v1/admin/export.pdf", headers=admin)
    assert pdf.status_code == 200 and pdf.content.startswith(b"%PDF")
    assert (await client.get("/api/v1/admin/export.csv", headers=admin)).status_code == 404


async def test_badge_crud_with_validation_and_audit(client, database):
    admin = await admin_headers(client)
    payload = {
        "code": code(),
        "name": "Uji Badge",
        "icon": "🧪",
        "description": "Badge uji",
        "criteria": {"type": "sessions_count", "gte": 3},
    }
    created = await client.post("/api/v1/admin/resources/badges", json=payload, headers=admin)
    assert created.status_code == 201, created.text
    badge = created.json()
    dup = await client.post("/api/v1/admin/resources/badges", json=payload, headers=admin)
    assert dup.status_code == 409
    bad = await client.post(
        "/api/v1/admin/resources/badges",
        json={**payload, "code": code(), "criteria": {"type": "magic", "gte": 1}},
        headers=admin,
    )
    assert bad.status_code == 422 and bad.json()["error"]["fields"]

    updated = await client.put(
        f"/api/v1/admin/resources/badges/{badge['id']}",
        json={**payload, "name": "Badge Baru"},
        headers=admin,
    )
    assert updated.json()["name"] == "Badge Baru"
    assert (
        await client.delete(f"/api/v1/admin/resources/badges/{badge['id']}", headers=admin)
    ).status_code == 204

    actions = [
        a["action"]
        async for a in database["audit_logs"]
        .find({"entity_id": ObjectId(badge["id"])})
        .sort("_id", 1)
    ]
    assert actions == ["badges.create", "badges.update", "badges.delete"]
    audit = (
        await client.get("/api/v1/admin/audit", params={"entity_type": "badges"}, headers=admin)
    ).json()
    assert audit["items"][0]["action"] == "badges.delete"


async def test_function_hierarchy(client, database):
    admin = await admin_headers(client)
    url = "/api/v1/admin/resources/functions"
    root = (
        await client.post(url, json={"name": "Direktorat", "code": code()}, headers=admin)
    ).json()
    child = (
        await client.post(
            url, json={"name": "Divisi", "code": code(), "parent_id": root["id"]}, headers=admin
        )
    ).json()
    grandchild = (
        await client.post(
            url, json={"name": "Tim", "code": code(), "parent_id": child["id"]}, headers=admin
        )
    ).json()
    assert grandchild["ancestors"] == [root["id"], child["id"]]

    # Pindahkan divisi ke root baru → ancestors cucu ikut berubah.
    new_root = (
        await client.post(url, json={"name": "Direktorat 2", "code": code()}, headers=admin)
    ).json()
    await client.put(
        f"{url}/{child['id']}",
        json={"name": "Divisi", "code": child["code"], "parent_id": new_root["id"]},
        headers=admin,
    )
    moved = await database["functions"].find_one({"_id": ObjectId(grandchild["id"])})
    assert moved["ancestors"] == [ObjectId(new_root["id"]), ObjectId(child["id"])]

    cycle = await client.put(
        f"{url}/{child['id']}",
        json={"name": "Divisi", "code": child["code"], "parent_id": grandchild["id"]},
        headers=admin,
    )
    assert cycle.status_code == 422
    in_use = await client.delete(f"{url}/{child['id']}", headers=admin)
    assert in_use.status_code == 409


async def test_roles_permission_matrix_and_lockout(client, database):
    admin = await admin_headers(client)
    url = "/api/v1/admin/resources/roles"
    catalog = (await client.get("/api/v1/admin/permissions", headers=admin)).json()
    assert any(p["code"] == "admin.dashboard.view" for p in catalog)

    role = (
        await client.post(
            url,
            json={
                "code": code(),
                "name": "Analis",
                "permission_codes": ["admin.dashboard.view", "leaderboard.view"],
            },
            headers=admin,
        )
    ).json()
    unknown = await client.post(
        url, json={"code": code(), "name": "X", "permission_codes": ["god.mode"]}, headers=admin
    )
    assert unknown.status_code == 422

    # User mendapat role baru → akses langsung berlaku (cache permission dibersihkan).
    member = await onboarded_user(client)
    member_id = (await me(client, member))["id"]
    assert (await client.get("/api/v1/admin/dashboard", headers=member)).status_code == 403
    await client.put(
        f"/api/v1/admin/users/{member_id}", json={"role_id": role["id"]}, headers=admin
    )
    assert (await client.get("/api/v1/admin/dashboard", headers=member)).status_code == 200
    assert (await client.delete(f"{url}/{role['id']}", headers=admin)).status_code == 409  # dipakai

    admin_role = await database["roles"].find_one({"code": "admin"})
    lockout = await client.put(
        f"{url}/{admin_role['_id']}",
        json={"code": "admin", "name": "Admin", "permission_codes": ["admin.dashboard.view"]},
        headers=admin,
    )
    assert lockout.status_code == 422
    assert (await client.delete(f"{url}/{admin_role['_id']}", headers=admin)).status_code == 409


async def test_point_rule_change_applies_immediately(client, fake_clock, database):
    admin = await admin_headers(client)
    rule = await database["point_rules"].find_one({"code": "post_feed"})
    url = f"/api/v1/admin/resources/point-rules/{rule['_id']}"
    try:
        response = await client.put(
            url, json={"name": rule["name"], "points": 15, "daily_cap_count": 3}, headers=admin
        )
        assert response.status_code == 200 and response.json()["code"] == "post_feed"
        await database["quests"].update_many({}, {"$set": {"is_active": False}})
        reader = await onboarded_user(client)
        book = await create_book(client, reader)
        result = await quick_post(client, reader, book["id"], note())
        awarded = {a["rule_code"]: a["points"] for a in result["points"]["awarded"]}
        assert awarded["post_feed"] == 15
        assert (
            await client.post("/api/v1/admin/resources/point-rules", json={}, headers=admin)
        ).status_code == 405
    finally:
        await database["point_rules"].update_one({"_id": rule["_id"]}, {"$set": {"points": 10}})
        await database["quests"].update_many({}, {"$set": {"is_active": True}})


async def test_settings_validation(client, database):
    admin = await admin_headers(client)
    listed = (await client.get("/api/v1/admin/settings", headers=admin)).json()
    assert any(s["key"] == "session.min_minutes" for s in listed)

    async def put(key, value):
        return await client.put(
            f"/api/v1/admin/settings/{key}", json={"value": value}, headers=admin
        )

    assert (await put("session.min_minutes", 3)).status_code == 422
    assert (await put("unknown.key", 1)).status_code == 404
    bad_order = await put(
        "authenticity.thresholds", {"active_reader": 0.2, "warming_up": 0.5, "observer": 0}
    )
    assert bad_order.status_code == 422
    assert (await put("team.timezone", "Mars/Olympus")).status_code == 422
    ok = await put(
        "authenticity.thresholds", {"active_reader": 0.6, "warming_up": 0.3, "observer": 0}
    )
    assert ok.status_code == 200 and ok.json()["value"]["active_reader"] == 0.6
    await put("authenticity.thresholds", {"active_reader": 0.5, "warming_up": 0.25, "observer": 0})
    assert await database["audit_logs"].count_documents({"action": "setting.update"}) >= 2


async def test_suspend_user_revokes_sessions(client):
    admin = await admin_headers(client)
    member = await onboarded_user(client)
    member_me = await me(client, member)
    found = (
        await client.get("/api/v1/admin/users", params={"q": member_me["phone"]}, headers=admin)
    ).json()
    assert found["total"] == 1 and found["items"][0]["id"] == member_me["id"]

    await client.put(
        f"/api/v1/admin/users/{member_me['id']}", json={"status": "suspended"}, headers=admin
    )
    assert (await client.get("/api/v1/me", headers=member)).status_code == 401
    login = await client.post(
        "/api/v1/auth/login", json={"phone": member_me["phone"], "pin": USER_PIN}
    )
    assert login.status_code == 403

    admin_id = (await me(client, admin))["id"]
    self_change = await client.put(
        f"/api/v1/admin/users/{admin_id}", json={"status": "suspended"}, headers=admin
    )
    assert self_change.status_code == 422


async def test_report_and_moderate_with_point_reversal(client, fake_clock, database):
    admin = await admin_headers(client)
    author, reporter = await onboarded_user(client), await onboarded_user(client)
    book = await create_book(client, author)
    post = (await quick_post(client, author, book["id"], note()))["post"]
    points_before = (await me(client, author))["stats"]["points_total"]
    assert points_before > 0

    url = f"/api/v1/posts/{post['id']}/report"
    assert (
        await client.post(url, json={"reason": "Spam berulang"}, headers=reporter)
    ).status_code == 204
    assert (
        await client.post(url, json={"reason": "Spam lagi"}, headers=reporter)
    ).status_code == 409
    assert (
        await client.post(url, json={"reason": "Punya sendiri"}, headers=author)
    ).status_code == 422

    queue = (await client.get("/api/v1/admin/moderation", headers=admin)).json()
    flagged = next(p for p in queue["posts"] if p["id"] == post["id"])
    assert flagged["reports"][0]["reason"] == "Spam berulang"

    hidden = await client.post(
        f"/api/v1/admin/moderation/posts/{post['id']}",
        json={"action": "hide", "reason": "spam", "reverse_points": True},
        headers=admin,
    )
    assert hidden.json()["status"] == "hidden" and hidden.json()["reversed_entries"] >= 2
    assert (await client.get(f"/api/v1/posts/{post['id']}", headers=reporter)).status_code == 404
    assert (await me(client, author))["stats"]["points_total"] == 0
    reversals = await database["points_ledger"].count_documents(
        {"user_id": ObjectId((await me(client, author))["id"]), "rule_code": "reversal"}
    )
    assert reversals >= 2

    restored = await client.post(
        f"/api/v1/admin/moderation/posts/{post['id']}", json={"action": "restore"}, headers=admin
    )
    assert restored.json()["status"] == "visible"
    assert (await client.get(f"/api/v1/posts/{post['id']}", headers=reporter)).status_code == 200


async def test_hide_comment_updates_count(client, fake_clock):
    admin = await admin_headers(client)
    author, other = await onboarded_user(client), await onboarded_user(client)
    book = await create_book(client, author)
    post = (await quick_post(client, author, book["id"], note()))["post"]
    comment = (
        await client.post(
            f"/api/v1/posts/{post['id']}/comments",
            json={"content": "Komentar kasar"},
            headers=other,
        )
    ).json()["comment"]
    report_url = f"/api/v1/posts/{post['id']}/comments/{comment['id']}/report"
    own = await client.post(report_url, json={"reason": "Sendiri"}, headers=other)
    assert own.status_code == 422
    reported = await client.post(report_url, json={"reason": "Tidak sopan"}, headers=author)
    assert reported.status_code == 204
    again = await client.post(report_url, json={"reason": "Tidak sopan"}, headers=author)
    assert again.status_code == 409
    flagged = (await client.get("/api/v1/admin/moderation", headers=admin)).json()
    item = next(c for c in flagged["comments"] if c["id"] == comment["id"])
    assert item["reports"][0]["reason"] == "Tidak sopan"
    await client.post(
        f"/api/v1/admin/moderation/comments/{comment['id']}", json={"action": "hide"}, headers=admin
    )
    detail = (await client.get(f"/api/v1/posts/{post['id']}", headers=author)).json()
    assert detail["counts"]["comments"] == 0
    assert (await client.get(f"/api/v1/posts/{post['id']}/comments", headers=author)).json() == []


async def test_admin_edits_book_and_denormalized_posts(client, fake_clock):
    admin = await admin_headers(client)
    reader = await onboarded_user(client)
    book = await create_book(client, reader, title=f"Judul Typo {uuid.uuid4().hex[:4]}")
    post = (await quick_post(client, reader, book["id"], note()))["post"]
    payload = {
        "title": "Judul Benar",
        "authors": ["Penulis Asli"],
        "category_id": book["category"]["id"],
    }
    updated = await client.put(f"/api/v1/admin/books/{book['id']}", json=payload, headers=admin)
    assert updated.status_code == 200 and updated.json()["title"] == "Judul Benar"
    detail = (await client.get(f"/api/v1/posts/{post['id']}", headers=reader)).json()
    assert detail["book"]["title"] == "Judul Benar"

    other = await create_book(client, reader)
    clash = await client.put(f"/api/v1/admin/books/{other['id']}", json=payload, headers=admin)
    assert clash.status_code == 409


async def test_admin_resets_pin_and_unlocks(client, database):
    admin = await admin_headers(client)
    member = await onboarded_user(client)
    member_me = await me(client, member)
    for _ in range(5):
        await client.post("/api/v1/auth/login", json={"phone": member_me["phone"], "pin": "000001"})
    locked = (
        await client.get("/api/v1/admin/users", params={"q": member_me["phone"]}, headers=admin)
    ).json()["items"][0]
    assert locked["locked"] is True

    weak = await client.put(
        f"/api/v1/admin/users/{member_me['id']}/pin", json={"pin": "111111"}, headers=admin
    )
    assert weak.status_code == 422
    reset = await client.put(
        f"/api/v1/admin/users/{member_me['id']}/pin", json={"pin": "482913"}, headers=admin
    )
    assert reset.status_code == 204
    assert (await client.get("/api/v1/me", headers=member)).status_code == 401  # sesi dicabut
    login = await client.post(
        "/api/v1/auth/login", json={"phone": member_me["phone"], "pin": "482913"}
    )
    assert login.status_code == 200
    log = await database["audit_logs"].find_one({"action": "user.reset_pin"})
    assert log and "password_hash" not in (log.get("after") or {})


async def test_notification_template_is_data_driven(client, fake_clock, database):
    admin = await admin_headers(client)
    template = await database["notification_templates"].find_one({"type": "comment"})
    url = f"/api/v1/admin/resources/notification-templates/{template['_id']}"
    try:
        await client.put(
            url, json={"title": "Ada komentar dari {actor}!", "body": "{excerpt}"}, headers=admin
        )
        author, other = await onboarded_user(client), await onboarded_user(client)
        book = await create_book(client, author)
        post = (await quick_post(client, author, book["id"], note()))["post"]
        await client.post(
            f"/api/v1/posts/{post['id']}/comments", json={"content": "Halo"}, headers=other
        )
        items = (await client.get("/api/v1/notifications", headers=author)).json()["items"]
        assert any(n["title"].startswith("Ada komentar dari") for n in items)
    finally:
        await database["notification_templates"].update_one(
            {"_id": template["_id"]},
            {"$set": {"title": template["title"], "body": template["body"]}},
        )


async def test_generic_resources_listing(client):
    admin = await admin_headers(client)
    for name in [
        "functions",
        "roles",
        "point-rules",
        "badges",
        "quests",
        "levels",
        "book-categories",
        "notification-templates",
    ]:
        response = await client.get(f"/api/v1/admin/resources/{name}", headers=admin)
        assert response.status_code == 200, name
        assert response.json()["items"], name
    assert (await client.get("/api/v1/admin/resources/nope", headers=admin)).status_code == 404
    _ = auth_header


async def test_lookups_for_admin_forms(client):
    admin = await admin_headers(client)
    data = (await client.get("/api/v1/admin/lookups", headers=admin)).json()
    assert {"roles", "functions", "categories", "users"} <= set(data)
    assert {r["code"] for r in data["roles"]} >= {"member", "team_lead", "admin"}
    assert all("password_hash" not in u and "id" in u for u in data["users"])
    assert (
        await client.get("/api/v1/admin/resources/invite-codes", headers=admin)
    ).status_code == 404
